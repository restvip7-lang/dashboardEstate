"""
Builds templates/dashboard-data-template.xlsx: the workbook the team fills in
(upload it to Google Drive and it opens as a Google Sheet).

Every tab maps onto a part of the dashboard's data contract (src/data/types.ts).
Row 1 holds the column headers the importer reads, row 2 a hint on the
format, row 3 a worked example; real data starts on row 4. Tab and column
names must not be renamed: the importer finds data by them.

Run: python3 scripts/build-data-template.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "templates" / "dashboard-data-template.xlsx"

FIRST_ROW = 4  # first data row
# Check formulas are pre-filled on this many rows (years of entries); formats
# and drop-downs cover the whole range. Copy the last row down to extend.
CALC_ROWS = 1000
BRAND_BLUE = "115E86"
BRAND_RED = "BF1E25"

# Tab colours by who fills the tab.
OWNER_COLOURS = {
    "Директор": BRAND_RED,
    "Руководитель продаж": BRAND_BLUE,
    "Бухгалтер": "2E8B57",
    "Строительство": "E07B00",
    "Отдел продаж застройщика": "6A4C93",
    "Справочник": "808080",
}

FONT = "Arial"
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
HINT_FONT = Font(name=FONT, italic=True, color="595959", size=9)
EXAMPLE_FONT = Font(name=FONT, italic=True, color="7F7F7F", size=10)
BODY_FONT = Font(name=FONT, size=10)
HEAD_FILL = PatternFill("solid", fgColor=BRAND_BLUE)
CALC_HEAD_FILL = PatternFill("solid", fgColor="4F6D80")
HINT_FILL = PatternFill("solid", fgColor="F2F2F2")
EXAMPLE_FILL = PatternFill("solid", fgColor="F7F7F7")
CALC_FILL = PatternFill("solid", fgColor="E8EEF2")
THIN = Side(style="thin", color="D0D7DE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

DATE_FMT = "DD.MM.YYYY"
MONTH_FMT = "MM.YYYY"
EUR_FMT = '#,##0 "€"'
PCT_FMT = '0"%"'

PROJECTS = [
    ("P01", "Blue Sunlight 2 Residence", "Аланья, Махмутлар"),
    ("P02", "Riva Port Villas Side", "Сиде"),
    ("P03", "Day One Residence (Вилла)", "Аланья, Тепе"),
    ("P04", "Exodus Panorama Residence", "Стамбул, Картал"),
    ("P05", "Prime Stone Residence", "Газипаша, Аланья"),
    ("P06", "Prime Garden Residence", "Аланья, Оба"),
    ("P07", "Prime Botanic Residence", "Аланья, Махмутлар"),
    ("P08", "Exodus Twins Residence", "Аланья, Махмутлар"),
    ("P09", "Blue Dream Residence", "Газипаша, Аланья"),
    ("P10", "Exodus Riverside Residence", "Аланья, Демирташ"),
    ("P11", "Blue Sunlight", "Аланья, Махмутлар"),
    ("P12", "Hayat Heaven Residence", "Аланья, Авсаллар"),
    ("P13", "Exodus Resort Comfort City", "Аланья, Махмутлар"),
    ("P14", "Exodus Green Hill Residence", "Стамбул, Картал"),
    ("P15", "Prime Loft Residence", "Аланья, Махмутлар"),
    ("P16", "Faralya Residence", "Аланья, Паяллар"),
    ("P17", "Villa Rabbit Hill", "Аланья, Бекташ"),
    ("P18", "Exodus Hill Residence", "Аланья, Махмутлар"),
    ("P19", "Exodus Nature Residence", "Аланья, Оба"),
    ("P20", "Exodus Premium Town", "Аланья, Каргыджак"),
    ("P21", "Exodus Aqua Deluxe Konakli", "Аланья, Конаклы"),
    ("P22", "Exodus Dreams Residence", "Аланья, Паяллар"),
]

# Stage name, default weight (% of overall readiness). Same order as the dashboard.
STAGES = [
    ("Фундамент", 15),
    ("Каркас", 30),
    ("Инженерия", 20),
    ("Фасад", 15),
    ("Отделка", 15),
    ("Благоустройство", 5),
]

COUNTRIES = [
    ("Турция", "TR"), ("Россия", "RU"), ("Германия", "DE"), ("Казахстан", "KZ"), ("Украина", "UA"),
    ("Великобритания", "GB"), ("Иран", "IR"), ("Польша", "PL"), ("Нидерланды", "NL"), ("Швеция", "SE"),
    ("Азербайджан", "AZ"), ("Беларусь", "BY"), ("Узбекистан", "UZ"), ("Кыргызстан", "KG"), ("Израиль", "IL"),
    ("Норвегия", "NO"), ("Финляндия", "FI"), ("Дания", "DK"), ("Франция", "FR"), ("Италия", "IT"),
    ("Австрия", "AT"), ("Швейцария", "CH"), ("США", "US"), ("ОАЭ", "AE"), ("Ирак", "IQ"),
    ("Саудовская Аравия", "SA"), ("Другая", "XX"),
]

SIGNAL_TOPICS = [
    "Просроченная комиссия",
    "Темп плана продаж",
    "Сроки проекта",
    "Отчёт стройки устарел",
    "Просрочка покупателей",
]

wb = Workbook()
wb.remove(wb.active)


def list_ref(sheet: str, col: str, last: int = 2000, first: int = FIRST_ROW) -> str:
    return f"='{sheet}'!${col}${first}:${col}${last}"


def add_table(ws, owner, columns, example, last_row, prefill=None):
    """columns: list of dicts {name, hint, width, fmt?, valid?, formula?, note?}."""
    ws.sheet_properties.tabColor = OWNER_COLOURS[owner]
    for i, col in enumerate(columns, start=1):
        letter = get_column_letter(i)
        is_calc = "formula" in col
        head = ws.cell(row=1, column=i, value=col["name"])
        head.font = HEAD_FONT
        head.fill = CALC_HEAD_FILL if is_calc else HEAD_FILL
        head.alignment = Alignment(wrap_text=True, vertical="center")
        head.border = BORDER
        if col.get("note"):
            head.comment = Comment(col["note"], "Шаблон")
        hint = ws.cell(row=2, column=i, value=("считается сама — не заполнять" if is_calc else col["hint"]))
        hint.font = HINT_FONT
        hint.fill = HINT_FILL
        hint.alignment = Alignment(wrap_text=True, vertical="top")
        hint.border = BORDER
        ex = ws.cell(row=3, column=i, value=example[i - 1] if i - 1 < len(example) else None)
        ex.font = EXAMPLE_FONT
        ex.fill = EXAMPLE_FILL
        ex.border = BORDER
        if col.get("fmt"):
            ex.number_format = col["fmt"]
        ws.column_dimensions[letter].width = col["width"]
        for r in range(FIRST_ROW, last_row + 1):
            c = ws.cell(row=r, column=i)
            c.font = BODY_FONT
            if col.get("fmt"):
                c.number_format = col["fmt"]
            if is_calc and r < FIRST_ROW + CALC_ROWS:
                c.value = col["formula"].format(r=r)
                c.fill = CALC_FILL
        if col.get("valid"):
            dv = col["valid"]()
            ws.add_data_validation(dv)
            dv.add(f"{letter}{FIRST_ROW}:{letter}{last_row}")
    if prefill:
        for r_off, row in enumerate(prefill):
            for c_off, value in enumerate(row):
                if value is not None:
                    ws.cell(row=FIRST_ROW + r_off, column=c_off + 1, value=value)
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 40
    ws.freeze_panes = "B4"


# ---- validation factories -------------------------------------------------


def v_list(items):
    def make():
        dv = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True)
        dv.error = "Выберите значение из списка"
        dv.errorTitle = "Недопустимое значение"
        return dv

    return make


def v_range(ref):
    def make():
        dv = DataValidation(type="list", formula1=ref, allow_blank=True)
        dv.error = "Выберите значение из списка"
        dv.errorTitle = "Нет в справочнике"
        return dv

    return make


def v_date():
    dv = DataValidation(type="date", operator="between", formula1="DATE(2018,1,1)", formula2="DATE(2035,12,31)", allow_blank=True)
    dv.error = "Введите дату в формате ДД.ММ.ГГГГ"
    dv.errorTitle = "Нужна дата"
    return dv


def v_money():
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    dv.error = "Введите сумму числом, без знака € и пробелов"
    dv.errorTitle = "Нужно число"
    return dv


def v_signed_money():
    dv = DataValidation(type="decimal", allow_blank=True, operator="between", formula1="-100000000", formula2="100000000")
    dv.error = "Введите сумму числом (минус — возврат)"
    dv.errorTitle = "Нужно число"
    return dv


def v_pct():
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="100", allow_blank=True)
    dv.error = "Введите число от 0 до 100"
    dv.errorTitle = "Нужен процент"
    return dv


def v_int():
    dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    dv.error = "Введите целое число"
    dv.errorTitle = "Нужно целое число"
    return dv


PROJECT_LIST = v_range(list_ref("Проекты", "B", 200))
MANAGER_LIST = v_range(list_ref("Менеджеры", "B", 200))
COUNTRY_LIST = v_range(f"='Списки'!$A$2:$A${len(COUNTRIES) + 1}")
STAGE_LIST = v_list([s for s, _ in STAGES])

# ---- Инструкция -------------------------------------------------------------

ins = wb.create_sheet("Инструкция")
ins.sheet_properties.tabColor = "000000"
ins.column_dimensions["A"].width = 26
ins.column_dimensions["B"].width = 30
ins.column_dimensions["C"].width = 22
ins.column_dimensions["D"].width = 70

ins["A1"] = "STAYPROPERTY · Данные для дашборда директора"
ins["A1"].font = Font(name=FONT, bold=True, size=16, color=BRAND_BLUE)
ins["A2"] = (
    "Каждый ответственный вносит только свои вкладки. Дашборд сам считает итоги, планы, просрочку, "
    "готовность и статусы — сюда вносятся только факты: дата, сумма, кто, какой объект."
)
ins["A2"].font = Font(name=FONT, size=10)
ins.merge_cells("A2:D2")
ins["A2"].alignment = Alignment(wrap_text=True)
ins.row_dimensions[2].height = 30

ins["A4"] = "Как заполнять"
ins["A4"].font = Font(name=FONT, bold=True, size=12)
rules = [
    ("Строка 1", "Названия колонок. Не переименовывайте колонки и вкладки — по ним дашборд находит данные."),
    ("Строка 2 (серая)", "Подсказка: что и в каком формате вносить."),
    ("Строка 3 (серая, курсив)", "Пример заполнения. Не удаляйте и не меняйте — дашборд её пропускает."),
    ("С 4-й строки", "Ваши данные. Одна строка = одно событие (сделка, платёж, отчёт). Пустые строки между данными не оставляйте."),
    ("Колонки тёмно-серого цвета", "Считаются автоматически (проверки, готовность). Не заполнять и не стирать. Формулы стоят в первых 1000 строках; если строки закончатся — скопируйте последнюю строку вниз."),
    ("Выпадающие списки", "Проекты, менеджеры, страны, статусы — выбирайте из списка, не печатайте вручную."),
    ("Даты", "Формат ДД.ММ.ГГГГ, например 26.09.2026. Месяц плана — первое число месяца: 01.09.2026."),
    ("Суммы", "В евро, числом без знака € и пробелов. Комиссия — без НДС (KDV) и за вычетом доли партнёров."),
    ("Другая валюта", "Если договор не в евро — заполните колонки «Валюта», «Сумма в валюте» и «Курс» (курс ЦБ Турции на дату сделки; правило пока предварительное)."),
    ("Исправления", "Ошибочную строку исправьте на месте. Отменённую сделку не удаляйте — поставьте дату отмены."),
]
for i, (a, d) in enumerate(rules, start=5):
    ins.cell(row=i, column=1, value=a).font = Font(name=FONT, bold=True, size=10)
    c = ins.cell(row=i, column=2, value=d)
    c.font = Font(name=FONT, size=10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ins.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
    ins.row_dimensions[i].height = 28

start = 5 + len(rules) + 1
ins.cell(row=start, column=1, value="Вкладки и ответственные").font = Font(name=FONT, bold=True, size=12)
tabs_info = [
    ("Сделки агентства", "Руководитель продаж", "в день сделки", "Закрытые сделки агентства: договор подписан и получен первый платёж сверх брони."),
    ("Оплаты комиссии", "Бухгалтер", "по выписке банка", "Деньги, фактически поступившие агентству по сделкам."),
    ("План агентства", "Директор", "раз в месяц", "План отдела продаж: комиссия и количество сделок."),
    ("План менеджеров", "Директор", "раз в месяц", "Личные планы менеджеров по комиссии."),
    ("Менеджеры", "Директор", "при найме и увольнении", "Штат отдела продаж с датами работы."),
    ("Проекты", "Строительство", "один раз, далее по согласованию", "22 проекта: статус, утверждённый срок сдачи, ответственный."),
    ("График этапов", "Строительство", "один раз, далее по согласованию", "Утверждённый график по этапам: даты начала и окончания, вес этапа."),
    ("Отчёты стройки", "Строительство", "раз в неделю", "Процент выполнения этапов, прогноз сдачи, ссылка на фото."),
    ("Шахматка", "Отдел продаж застройщика", "при изменениях", "Квартиры и виллы: статус и цена по прайсу."),
    ("Продажи застройщика", "Отдел продаж застройщика", "при подписании договора", "Договоры покупателей и канал продажи."),
    ("График платежей", "Бухгалтер", "при подписании договора", "Сроки и суммы платежей покупателей по договору."),
    ("Платежи покупателей", "Бухгалтер", "по выписке банка", "Поступления от покупателей (минус — возврат)."),
    ("План застройщика", "Директор", "раз в месяц", "План продаж застройщика и ожидаемые первые взносы."),
    ("Внимание директора", "Директор", "по мере появления", "Кто отвечает за проблему и что делает. Сами проблемы дашборд находит сам."),
    ("Списки", "Справочник", "не менять", "Страны и прочие списки для выпадающих меню."),
]
head_row = start + 1
for j, h in enumerate(["Вкладка", "Кто заполняет", "Когда", "Что вносится"], start=1):
    c = ins.cell(row=head_row, column=j, value=h)
    c.font = HEAD_FONT
    c.fill = HEAD_FILL
    c.border = BORDER
for i, (tab, owner, when, what) in enumerate(tabs_info, start=head_row + 1):
    values = [tab, owner, when, what]
    for j, v in enumerate(values, start=1):
        c = ins.cell(row=i, column=j, value=v)
        c.font = Font(name=FONT, size=10, bold=(j == 1), color=(OWNER_COLOURS[owner] if j == 2 else None))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = BORDER
    ins.row_dimensions[i].height = 30

note_row = head_row + len(tabs_info) + 2
ins.cell(row=note_row, column=1, value="Важно").font = Font(name=FONT, bold=True, size=12, color=BRAND_RED)
c = ins.cell(
    row=note_row + 1,
    column=1,
    value=(
        "Таблица содержит коммерческие данные компании. Доступ к ней — только ответственным. "
        "Не публикуйте таблицу в интернет и не выкладывайте её в публичный репозиторий."
    ),
)
c.font = Font(name=FONT, size=10)
c.alignment = Alignment(wrap_text=True)
ins.merge_cells(start_row=note_row + 1, start_column=1, end_row=note_row + 1, end_column=4)
ins.row_dimensions[note_row + 1].height = 30

# ---- Сделки агентства -------------------------------------------------------

ws = wb.create_sheet("Сделки агентства")
add_table(
    ws,
    "Руководитель продаж",
    [
        {"name": "ID сделки", "hint": "номер сделки в Битрикс24 или свой, без повторов", "width": 14},
        {"name": "Дата закрытия", "hint": "договор подписан и получен первый платёж", "width": 13, "fmt": DATE_FMT, "valid": v_date},
        {"name": "Дата отмены", "hint": "только если сделка отменена", "width": 12, "fmt": DATE_FMT, "valid": v_date},
        {"name": "Тип объекта", "hint": "новостройка или вторичка", "width": 13, "valid": v_list(["Новостройка", "Вторичка"])},
        {"name": "Наш проект", "hint": "Да — объект одного из 22 проектов компании", "width": 10, "valid": v_list(["Да", "Нет"])},
        {"name": "Проект", "hint": "если «Наш проект» = Да", "width": 26, "valid": PROJECT_LIST},
        {"name": "Объект", "hint": "блок, номер квартиры или адрес", "width": 20},
        {"name": "Цена, €", "hint": "цена по договору", "width": 12, "fmt": EUR_FMT, "valid": v_money},
        {"name": "Комиссия агентства, €", "hint": "без НДС и за вычетом доли партнёров", "width": 14, "fmt": EUR_FMT, "valid": v_money},
        {"name": "Менеджер 1", "hint": "основной менеджер", "width": 18, "valid": MANAGER_LIST},
        {"name": "Доля 1, %", "hint": "100, если менеджер один", "width": 9, "fmt": PCT_FMT, "valid": v_pct},
        {"name": "Менеджер 2", "hint": "если сделку вели двое", "width": 18, "valid": MANAGER_LIST},
        {"name": "Доля 2, %", "hint": "доли в сумме = 100", "width": 9, "fmt": PCT_FMT, "valid": v_pct},
        {"name": "Страна покупателя", "hint": "гражданство основного покупателя", "width": 16, "valid": COUNTRY_LIST},
        {"name": "Срок оплаты комиссии 1", "hint": "когда должны заплатить", "width": 13, "fmt": DATE_FMT, "valid": v_date},
        {"name": "Сумма 1, €", "hint": "", "width": 11, "fmt": EUR_FMT, "valid": v_money},
        {"name": "Срок оплаты комиссии 2", "hint": "если застройщик платит частями", "width": 13, "fmt": DATE_FMT, "valid": v_date},
        {"name": "Сумма 2, €", "hint": "суммы 1+2 = комиссия", "width": 11, "fmt": EUR_FMT, "valid": v_money},
        {"name": "Валюта", "hint": "если договор не в евро", "width": 8, "valid": v_list(["EUR", "USD", "TRY", "GBP", "RUB"])},
        {"name": "Сумма в валюте", "hint": "цена в валюте договора", "width": 12, "valid": v_money},
        {"name": "Курс", "hint": "евро за 1 единицу валюты", "width": 9},
        {"name": "Комментарий", "hint": "", "width": 24},
        {
            "name": "Проверка",
            "hint": "",
            "width": 22,
            "formula": '=IF(A{r}="","",IF(ABS(K{r}+M{r}-100)>0.01,"Доли ≠ 100%",IF(ABS(P{r}+R{r}-I{r})>1,"Сумма графика ≠ комиссии",IF(AND(E{r}="Да",F{r}=""),"Укажите проект","OK"))))',
            "note": "Проверяет, что доли менеджеров дают 100%, график оплаты комиссии равен комиссии, а у нашего объекта указан проект.",
        },
    ],
    ["D-0001", "26.09.2026", None, "Новостройка", "Да", "Exodus Twins Residence", "Блок B, кв. 12", 215000, 10750,
     "Анна Волкова", 100, None, None, "Россия", "10.10.2026", 5375, "10.12.2026", 5375, None, None, None, "застройщик платит 2 частями", "OK"],
    last_row=2000,
)

# ---- Оплаты комиссии --------------------------------------------------------

ws = wb.create_sheet("Оплаты комиссии")
add_table(
    ws,
    "Бухгалтер",
    [
        {"name": "ID сделки", "hint": "из вкладки «Сделки агентства»", "width": 14, "valid": v_range(list_ref("Сделки агентства", "A"))},
        {"name": "Дата поступления", "hint": "дата в выписке банка", "width": 14, "fmt": DATE_FMT, "valid": v_date},
        {"name": "Сумма, €", "hint": "минус — возврат", "width": 12, "fmt": EUR_FMT, "valid": v_signed_money},
        {"name": "Комментарий", "hint": "номер платёжки, плательщик", "width": 28},
        {
            "name": "Проверка",
            "hint": "",
            "width": 20,
            "formula": "=IF(A{r}=\"\",\"\",IF(COUNTIF('Сделки агентства'!$A$4:$A$2000,A{r})>0,\"OK\",\"Нет такой сделки\"))",
            "note": "Проверяет, что сделка с таким ID есть на вкладке «Сделки агентства».",
        },
    ],
    ["D-0001", "12.10.2026", 5375, "п/п 1452, Exodus", "OK"],
    last_row=3000,
)

# ---- Планы ------------------------------------------------------------------

ws = wb.create_sheet("План агентства")
add_table(
    ws,
    "Директор",
    [
        {"name": "Месяц", "hint": "первое число месяца: 01.09.2026", "width": 12, "fmt": MONTH_FMT, "valid": v_date},
        {"name": "План комиссии, €", "hint": "начисленная комиссия отдела", "width": 16, "fmt": EUR_FMT, "valid": v_money},
        {"name": "План сделок", "hint": "количество закрытых сделок", "width": 12, "valid": v_int},
    ],
    ["01.09.2026", 300000, 30],
    last_row=200,
)

ws = wb.create_sheet("План менеджеров")
add_table(
    ws,
    "Директор",
    [
        {"name": "Месяц", "hint": "первое число месяца", "width": 12, "fmt": MONTH_FMT, "valid": v_date},
        {"name": "Менеджер", "hint": "из вкладки «Менеджеры»", "width": 20, "valid": MANAGER_LIST},
        {"name": "План комиссии, €", "hint": "личный план", "width": 16, "fmt": EUR_FMT, "valid": v_money},
    ],
    ["01.09.2026", "Анна Волкова", 45000],
    last_row=2000,
)

ws = wb.create_sheet("Менеджеры")
add_table(
    ws,
    "Директор",
    [
        {"name": "ID", "hint": "M01, M02…", "width": 8},
        {"name": "Имя", "hint": "как показывать на экране", "width": 22},
        {"name": "Отдел", "hint": "", "width": 12, "valid": v_list(["Продажи"])},
        {"name": "Дата начала", "hint": "первый рабочий день", "width": 13, "fmt": DATE_FMT, "valid": v_date},
        {"name": "Дата окончания", "hint": "только при увольнении", "width": 13, "fmt": DATE_FMT, "valid": v_date},
    ],
    ["M01", "Анна Волкова", "Продажи", "01.03.2022", None],
    last_row=200,
)

# ---- Стройка ----------------------------------------------------------------

ws = wb.create_sheet("Проекты")
add_table(
    ws,
    "Строительство",
    [
        {"name": "ID", "hint": "не менять", "width": 7},
        {"name": "Проект", "hint": "название, как в реестре", "width": 30},
        {"name": "Локация", "hint": "город, район", "width": 22},
        {"name": "Статус", "hint": "стадия проекта", "width": 14, "valid": v_list(["Предпродажи", "Строительство", "Сдан"])},
        {"name": "Утверждённый срок сдачи", "hint": "передача ключей; прогноз — во вкладке отчётов", "width": 15, "fmt": DATE_FMT, "valid": v_date},
        {"name": "Ответственный", "hint": "руководитель проекта", "width": 20},
    ],
    ["P00", "Название проекта", "Аланья, Махмутлар", "Строительство", "31.12.2026", "Имя Фамилия"],
    last_row=200,
    prefill=[(pid, name, loc) for pid, name, loc in PROJECTS],
)

ws = wb.create_sheet("График этапов")
add_table(
    ws,
    "Строительство",
    [
        {"name": "Проект", "hint": "", "width": 30, "valid": PROJECT_LIST},
        {"name": "Этап", "hint": "", "width": 16, "valid": STAGE_LIST},
        {"name": "Вес, %", "hint": "доля этапа в готовности; по проекту = 100", "width": 9, "fmt": PCT_FMT, "valid": v_pct},
        {"name": "Начало по плану", "hint": "утверждённый график", "width": 14, "fmt": DATE_FMT, "valid": v_date},
        {"name": "Окончание по плану", "hint": "утверждённый график", "width": 14, "fmt": DATE_FMT, "valid": v_date},
        {
            "name": "Сумма весов проекта",
            "hint": "",
            "width": 14,
            "formula": '=IF(A{r}="","",SUMIF($A$4:$A$400,A{r},$C$4:$C$400))',
            "note": "По каждому проекту веса этапов должны давать 100.",
        },
    ],
    ["Exodus Twins Residence", "Каркас", 30, "01.06.2025", "15.02.2026", 100],
    last_row=400,
    prefill=[(name, stage, weight) for _, name, _ in PROJECTS for stage, weight in STAGES],
)

stage_cols = ["C", "D", "E", "F", "G", "H"]
readiness_terms = "+".join(
    f"{col}{{r}}*SUMIFS('График этапов'!$C$4:$C$400,'График этапов'!$A$4:$A$400,B{{r}},'График этапов'!$B$4:$B$400,\"{stage}\")"
    for col, (stage, _) in zip(stage_cols, STAGES)
)
ws = wb.create_sheet("Отчёты стройки")
add_table(
    ws,
    "Строительство",
    [
        {"name": "Дата отчёта", "hint": "раз в неделю", "width": 12, "fmt": DATE_FMT, "valid": v_date},
        {"name": "Проект", "hint": "", "width": 28, "valid": PROJECT_LIST},
        *[
            {"name": f"{stage}, %", "hint": "0–100 от объёма этапа", "width": 11, "fmt": PCT_FMT, "valid": v_pct}
            for stage, _ in STAGES
        ],
        {"name": "Прогноз сдачи", "hint": "актуальная оценка передачи ключей", "width": 13, "fmt": DATE_FMT, "valid": v_date},
        {"name": "Ссылка на фото", "hint": "папка с фото на дату отчёта", "width": 26},
        {"name": "Комментарий", "hint": "риски, причины отставания", "width": 30},
        {
            "name": "Готовность, %",
            "hint": "",
            "width": 12,
            "fmt": PCT_FMT,
            "formula": "=IF(B{r}=\"\",\"\",IF(SUMIF('График этапов'!$A$4:$A$400,B{r},'График этапов'!$C$4:$C$400)=0,\"нет графика\","
            + f"({readiness_terms})/SUMIF('График этапов'!$A$4:$A$400,B{{r}},'График этапов'!$C$4:$C$400)))",
            "note": "Взвешенная готовность по весам этапов из вкладки «График этапов».",
        },
    ],
    ["25.09.2026", "Exodus Twins Residence", 100, 100, 90, 70, 20, 0, "31.12.2026", "https://drive.google.com/…", "фасад по графику", 76.5],
    last_row=1500,
)

# ---- Застройщик: продажи и деньги --------------------------------------------

ws = wb.create_sheet("Шахматка")
add_table(
    ws,
    "Отдел продаж застройщика",
    [
        {"name": "ID юнита", "hint": "проект-блок-номер, без повторов", "width": 16},
        {"name": "Проект", "hint": "", "width": 28, "valid": PROJECT_LIST},
        {"name": "Блок / номер", "hint": "", "width": 12},
        {"name": "Площадь, м²", "hint": "", "width": 11, "valid": v_money},
        {"name": "Статус", "hint": "текущий статус", "width": 14, "valid": v_list(["Продано", "Бронь", "Свободно", "Не в продаже"])},
        {"name": "Цена по прайсу, €", "hint": "действующий прайс", "width": 15, "fmt": EUR_FMT, "valid": v_money},
    ],
    ["P08-B-012", "Exodus Twins Residence", "B-12", 78, "Продано", 220000],
    last_row=5000,
)

ws = wb.create_sheet("Продажи застройщика")
add_table(
    ws,
    "Отдел продаж застройщика",
    [
        {"name": "ID договора", "hint": "номер договора, без повторов", "width": 14},
        {"name": "Дата договора", "hint": "договор подписан и получен первый платёж", "width": 13, "fmt": DATE_FMT, "valid": v_date},
        {"name": "Дата расторжения", "hint": "только при расторжении", "width": 13, "fmt": DATE_FMT, "valid": v_date},
        {"name": "Проект", "hint": "", "width": 28, "valid": PROJECT_LIST},
        {"name": "ID юнита", "hint": "из «Шахматки»", "width": 16, "valid": v_range(list_ref("Шахматка", "A", 5000))},
        {"name": "Цена договора, €", "hint": "", "width": 14, "fmt": EUR_FMT, "valid": v_money},
        {"name": "Канал", "hint": "кто продал", "width": 16, "valid": v_list(["Наше агентство", "Партнёр", "Прямая продажа"])},
        {"name": "ID сделки агентства", "hint": "если продало наше агентство", "width": 14},
        {
            "name": "Проверка",
            "hint": "",
            "width": 22,
            "formula": "=IF(A{r}=\"\",\"\",IF(ABS(SUMIF('График платежей'!$A$4:$A$5000,A{r},'График платежей'!$C$4:$C$5000)-F{r})>1,\"График платежей ≠ цене\",\"OK\"))",
            "note": "Сумма платежей по графику должна равняться цене договора.",
        },
    ],
    ["C-2026-081", "26.09.2026", None, "Exodus Twins Residence", "P08-B-012", 215000, "Наше агентство", "D-0001", "OK"],
    last_row=5000,
)

ws = wb.create_sheet("График платежей")
add_table(
    ws,
    "Бухгалтер",
    [
        {"name": "ID договора", "hint": "из «Продаж застройщика»", "width": 14, "valid": v_range(list_ref("Продажи застройщика", "A", 5000))},
        {"name": "Срок платежа", "hint": "по договору", "width": 13, "fmt": DATE_FMT, "valid": v_date},
        {"name": "Сумма, €", "hint": "одна строка = один платёж графика", "width": 12, "fmt": EUR_FMT, "valid": v_money},
        {
            "name": "Проверка",
            "hint": "",
            "width": 20,
            "formula": "=IF(A{r}=\"\",\"\",IF(COUNTIF('Продажи застройщика'!$A$4:$A$5000,A{r})>0,\"OK\",\"Нет такого договора\"))",
        },
    ],
    ["C-2026-081", "29.09.2026", 75250, "OK"],
    last_row=5000,
)

ws = wb.create_sheet("Платежи покупателей")
add_table(
    ws,
    "Бухгалтер",
    [
        {"name": "ID договора", "hint": "из «Продаж застройщика»", "width": 14, "valid": v_range(list_ref("Продажи застройщика", "A", 5000))},
        {"name": "Дата поступления", "hint": "дата в выписке банка", "width": 14, "fmt": DATE_FMT, "valid": v_date},
        {"name": "Сумма, €", "hint": "минус — возврат", "width": 12, "fmt": EUR_FMT, "valid": v_signed_money},
        {"name": "Комментарий", "hint": "номер платёжки", "width": 26},
        {
            "name": "Проверка",
            "hint": "",
            "width": 20,
            "formula": "=IF(A{r}=\"\",\"\",IF(COUNTIF('Продажи застройщика'!$A$4:$A$5000,A{r})>0,\"OK\",\"Нет такого договора\"))",
        },
    ],
    ["C-2026-081", "30.09.2026", 75250, "п/п 2201", "OK"],
    last_row=5000,
)

ws = wb.create_sheet("План застройщика")
add_table(
    ws,
    "Директор",
    [
        {"name": "Месяц", "hint": "первое число месяца", "width": 12, "fmt": MONTH_FMT, "valid": v_date},
        {"name": "План продаж, шт.", "hint": "", "width": 13, "valid": v_int},
        {"name": "План продаж, €", "hint": "", "width": 15, "fmt": EUR_FMT, "valid": v_money},
        {"name": "План первых взносов, €", "hint": "ожидаемые первые платежи по новым продажам", "width": 18, "fmt": EUR_FMT, "valid": v_money},
    ],
    ["01.09.2026", 30, 6000000, 2600000],
    last_row=200,
)

# ---- Внимание директора -----------------------------------------------------

ws = wb.create_sheet("Внимание директора")
add_table(
    ws,
    "Директор",
    [
        {"name": "Тема", "hint": "какую проблему закрываем", "width": 24, "valid": v_list(SIGNAL_TOPICS)},
        {"name": "Проект", "hint": "для тем «Сроки проекта» и «Отчёт стройки»", "width": 28, "valid": PROJECT_LIST},
        {"name": "Ответственный", "hint": "кто решает", "width": 20},
        {"name": "Действие", "hint": "коротко, будет на экране", "width": 36},
        {"name": "Срок", "hint": "", "width": 12, "fmt": DATE_FMT, "valid": v_date},
    ],
    ["Сроки проекта", "Exodus Panorama Residence", "Бурак Шахин", "усилить график подрядчика", "05.10.2026"],
    last_row=300,
)

# ---- Списки -----------------------------------------------------------------

ws = wb.create_sheet("Списки")
ws.sheet_properties.tabColor = OWNER_COLOURS["Справочник"]
for j, h in enumerate(["Страна", "Код"], start=1):
    c = ws.cell(row=1, column=j, value=h)
    c.font = HEAD_FONT
    c.fill = HEAD_FILL
for i, (name, code) in enumerate(COUNTRIES, start=2):
    ws.cell(row=i, column=1, value=name).font = BODY_FONT
    ws.cell(row=i, column=2, value=code).font = BODY_FONT
ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 8

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"saved {OUT}")
